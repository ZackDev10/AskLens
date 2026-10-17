import io

import pytest
from fastapi.testclient import TestClient

from app.api import deps
from app.ingestion.pipeline import IngestionService
from app.main import app
from tests.fakes import FakeEmbedder, FakeRepo, make_settings


@pytest.fixture
def repo():
    return FakeRepo()


@pytest.fixture
def embedder():
    return FakeEmbedder()


@pytest.fixture
def settings():
    return make_settings()


@pytest.fixture
def service(repo, embedder, settings):
    return IngestionService(repo, embedder, settings)


@pytest.fixture
def client(repo, embedder, settings):
    app.dependency_overrides[deps.get_repo] = lambda: repo
    app.dependency_overrides[deps.get_embedder] = lambda: embedder
    app.dependency_overrides[deps.get_settings] = lambda: settings
    yield TestClient(app)
    app.dependency_overrides.clear()


def make_pdf(pages: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for text in pages:
        y = 800
        for line in text.split("\n"):
            c.drawString(60, y, line)
            y -= 16
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_heading("Warranty", level=1)
    d.add_paragraph("All electrical items carry a twelve month warranty.")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "Item"
    t.rows[0].cells[1].text = "Period"
    t.rows[1].cells[0].text = "Desk lamp"
    t.rows[1].cells[1].text = "12 months"
    d.add_paragraph("Contact support to start a claim.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
